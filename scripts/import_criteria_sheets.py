"""
Импорт оценок из рабочей таблицы «по листу на сотрудника»
(напр. «Черновик.xlsx»: лист «оп2. Попова», критерии в строках, звонки в колонках).

Формат листа:
  строка 2, колонки E..  — ссылка на сделку Битрикс24 (одна колонка = один звонок)
  строка 3               — пометка «на калибровку» над колонкой (необязательно)
  колонка B              — текст критерия (строки с числовым весом в колонке D)
  колонки E..            — ответ: «Да» / «Нет» / «Не мог» (= Н/П); примечание к ячейке → комментарий

Чек-лист определяется по совпадению текстов критериев. Сотрудник,
дата звонка (дата презентации, иначе DATE_CREATE), стадия и категория —
из Битрикс по сделке, как при ручном создании оценки;
отдел — из названия листа («оп3. …» → «Отдел продления 3»). Листы «Шаблон*» пропускаются.

Запуск:
  .venv/bin/python scripts/import_criteria_sheets.py Черновик.xlsx --evaluator <логин> --eval-date ДД.ММ.ГГГГ --dry-run
"""
from __future__ import annotations
import sys
import os
import re
import json
import argparse
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import selectinload

from app.bitrix import get_deal
from app.database import SessionLocal
from app.models import Checklist, Block, Evaluation, EvaluationItem, User
from app.scoring import calculate_scores, MONTH_NAMES

_VALUES = {"да": "yes", "нет": "no", "не мог": "na", "н/п": "na"}
_DEAL_RE = re.compile(r"/crm/deal/details/(\d+)")
_SHEET_DEPT_RE = re.compile(r"^\s*оп\s*(\d+)", re.I)
FIRST_COL = 5  # E


def _norm(text) -> str:
    return " ".join(str(text or "").replace('""', '"').split())


def _operator_name(bitrix_name: str) -> str:
    # Битрикс: «Мария (Отпуск с … по …) Чекмезова» → «Чекмезова Мария», как в остальных оценках
    parts = re.sub(r"\([^)]*\)", " ", bitrix_name).split()
    return " ".join([parts[-1]] + parts[:-1]) if len(parts) >= 2 else " ".join(parts)


def _pick_checklist(texts: list[str], checklists: list[Checklist]):
    for cl in checklists:
        by_text = {_norm(c.text): c for b in cl.blocks for c in b.criteria}
        if len(by_text) == len(texts) and all(_norm(t) in by_text for t in texts):
            return cl, by_text
    return None, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--evaluator", required=True, help="логин оценщика (users.username)")
    ap.add_argument("--eval-date", required=True, help="дата оценки ДД.ММ.ГГГГ")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    db = SessionLocal()
    evaluator = db.query(User).filter(User.username == args.evaluator).first()
    if not evaluator:
        raise SystemExit(f"Пользователь {args.evaluator} не найден")
    eval_dt = datetime.strptime(args.eval_date, "%d.%m.%Y")
    checklists = db.query(Checklist).options(
        selectinload(Checklist.blocks).selectinload(Block.criteria)
    ).all()
    existing = {(e.deal_id, e.checklist_id) for e in db.query(Evaluation).all()}

    wb = openpyxl.load_workbook(args.file)
    to_add: list[Evaluation] = []
    errors = 0
    for ws in wb.worksheets:
        if ws.title.lower().startswith("шаблон"):
            continue
        crit_rows = [r for r in range(6, ws.max_row + 1)
                     if ws.cell(r, 2).value and isinstance(ws.cell(r, 4).value, (int, float))]
        cl, by_text = _pick_checklist([ws.cell(r, 2).value for r in crit_rows], checklists)
        if not cl:
            print(f"ОШИБКА лист «{ws.title}»: критерии не совпадают ни с одним чек-листом")
            errors += 1
            continue
        # отдел — из названия листа («оп3. Чекмезова»): в Битрикс сейчас может стоять уже другой
        m_dept = _SHEET_DEPT_RE.match(ws.title)
        sheet_dept = f"Отдел продления {m_dept.group(1)}" if m_dept else None
        print(f"Лист «{ws.title}» → чек-лист {cl.id} «{cl.name}», {sheet_dept or 'отдел из Битрикс'}")

        col = FIRST_COL
        while ws.cell(2, col).value:
            m = _DEAL_RE.search(str(ws.cell(2, col).value))
            letter = get_column_letter(col)
            col += 1
            if not m:
                print(f"  ОШИБКА {letter}2: не ссылка на сделку")
                errors += 1
                continue
            deal_id = m.group(1)
            calibration = "калибр" in str(ws.cell(3, col - 1).value or "").lower()

            items = []
            for r in crit_rows:
                cell = ws.cell(r, col - 1)
                raw = str(cell.value or "").strip().lower()
                v = _VALUES.get(raw)
                if v is None:
                    if raw:
                        print(f"  ОШИБКА {cell.coordinate}: непонятный ответ «{cell.value}»")
                        errors += 1
                    continue
                comment = None
                if cell.comment and cell.comment.text.strip():
                    comment = json.dumps(
                        [{"text": cell.comment.text.strip(), "flag": "", "time": ""}],
                        ensure_ascii=False)
                items.append((by_text[_norm(ws.cell(r, 2).value)].id, v, comment))

            if (deal_id, cl.id) in existing:
                print(f"  {letter}: сделка {deal_id} уже есть в базе по этому чек-листу — пропуск")
                continue
            existing.add((deal_id, cl.id))

            deal = get_deal(deal_id)
            if not deal:
                print(f"  ОШИБКА {letter}: сделка {deal_id} не найдена в Битрикс")
                errors += 1
                continue
            call_dt = deal.presentation_date or deal.deal_date
            if call_dt and not isinstance(call_dt, datetime):
                call_dt = datetime.combine(call_dt, datetime.min.time())
            operator = _operator_name(deal.operator_name)
            department = sheet_dept or deal.department
            total, _ = calculate_scores(items, cl)
            n_comm = sum(1 for i in items if i[2])
            print(f"  {letter}: сделка {deal_id} · {operator} · {department} · "
                  f"звонок {call_dt:%d.%m.%Y} · {deal.stage} · итог {total}% · "
                  f"ответов {len(items)}, комм. {n_comm}{' · калибровка' if calibration else ''}")

            ev = Evaluation(
                checklist_id=cl.id,
                deal_id=deal_id,
                deal_url=f"https://entera.bitrix24.ru/crm/deal/details/{deal_id}/",
                operator_name=operator,
                eval_date=call_dt,
                week_num=call_dt.isocalendar()[1] if call_dt else None,
                week_year=call_dt.year if call_dt else None,
                month=MONTH_NAMES[call_dt.month - 1] if call_dt else None,
                department=department,
                stage=deal.stage,
                client_category=deal.client_category,
                total_score=total,
                evaluator_id=evaluator.id,
                status="published",
                created_at=eval_dt,
                is_calibration=calibration,
            )
            ev.items = [EvaluationItem(criterion_id=c, value=v, comment=cm) for c, v, cm in items]
            to_add.append(ev)

    print(f"\nИтого к импорту: {len(to_add)} оценок, ошибок: {errors}")
    if errors:
        raise SystemExit("Есть ошибки — импорт не выполнен")
    if args.dry_run:
        print("--dry-run: в базу ничего не записано")
        return
    db.add_all(to_add)
    db.commit()
    print("Записано в базу.")


if __name__ == "__main__":
    main()
