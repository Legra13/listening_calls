"""
Импорт оценок из xlsx-выгрузок отчёта «Результаты сотрудников»
(лист «По оценкам», разрез «детально по критериям»).

Используется для восстановления данных, потерянных вместе со старым сервером.
Чек-лист определяется по совпадению заголовков столбцов с текстами критериев.
Итоговый балл пересчитывается через scoring.calculate_scores и сверяется
с колонкой «Итог %» из файла — расхождение > 0.15 п.п. выводится как ошибка.

Пропускаются: файлы в разрезе «группы» (без критериев) и оценки, которые
уже есть в базе (ключ: сделка + чек-лист + дата звонка + оператор + итог).

Запуск:
  .venv/bin/python scripts/import_employee_xlsx.py <файлы.xlsx...> --dry-run
  .venv/bin/python scripts/import_employee_xlsx.py <файлы.xlsx...>
"""
from __future__ import annotations
import sys
import os
import json
import argparse
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from sqlalchemy.orm import selectinload

from app.database import SessionLocal
from app.models import Checklist, Block, Evaluation, EvaluationItem, User
from app.scoring import calculate_scores, MONTH_NAMES

META_COLS = 5   # Сотрудник, Отдел, Дата звонка, Дата оценки, Оценивал
TAIL_COLS = 4   # Итог %, Сделка, Стадия, Комментарий

_VALUES = {"Да": "yes", "Нет": "no", "Н/П": "na"}
_STAGES = {"Успешна": "сделка успешна", "Не продал": "не смог продать", "В работе": "в работе"}


def _norm(text: str) -> str:
    # в текстах критериев в БД кавычки местами удвоены ("" вместо ")
    return " ".join((text or "").replace('""', '"').split())


def _date(v) -> datetime | None:
    if isinstance(v, datetime):
        return v
    if not v or v == "—":
        return None
    return datetime.strptime(str(v).strip(), "%d.%m.%Y")


def _pick_checklist(headers: list[str], checklists: list[Checklist]):
    best = None
    for cl in checklists:
        by_text = {_norm(c.text): c for b in cl.blocks for c in b.criteria}
        if len(by_text) != len(headers):
            continue
        if all(_norm(h) in by_text for h in headers):
            if best:
                raise SystemExit(f"Неоднозначно: подходят чек-листы {best[0].id} и {cl.id}")
            best = (cl, [by_text[_norm(h)] for h in headers])
    return best


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    db = SessionLocal()
    checklists = db.query(Checklist).options(
        selectinload(Checklist.blocks).selectinload(Block.criteria)
    ).all()
    users = {(u.full_name or u.username): u for u in db.query(User).all()}

    existing = {
        (e.deal_id, e.checklist_id, e.eval_date.date() if e.eval_date else None,
         e.operator_name, round(e.total_score or 0, 1))
        for e in db.query(Evaluation).all()
    }

    to_add: list[Evaluation] = []
    errors = 0
    for path in args.files:
        name = os.path.basename(path)
        ws = openpyxl.load_workbook(path)["По оценкам"]
        header_cells = list(ws[1])
        crit_headers = [(c.value or "").strip() for c in header_cells[META_COLS:-TAIL_COLS]]
        picked = _pick_checklist(crit_headers, checklists)
        if not picked:
            print(f"ПРОПУСК {name}: столбцы не совпадают ни с одним чек-листом "
                  f"(вероятно, разрез «группы» — нет ответов по критериям)")
            continue
        cl, crits = picked
        added = dup = 0
        for r in ws.iter_rows(min_row=2):
            vals = [c.value for c in r]
            if not vals[0]:
                continue
            operator, dept, call_d, eval_d, evaluator = vals[:META_COLS]
            total_file, deal_cell, stage_label, general = vals[-TAIL_COLS:]
            deal_id = str(deal_cell or "").lstrip("#").strip() or None
            call_dt, eval_dt = _date(call_d), _date(eval_d)

            items = []
            for i, crit in enumerate(crits):
                cell = r[META_COLS + i]
                v = _VALUES.get((cell.value or "").strip())
                if v is None:
                    continue  # «—» — ответа нет, при расчёте = Н/П
                comment = None
                if cell.comment and cell.comment.text.strip():
                    comment = json.dumps(
                        [{"text": cell.comment.text.strip(), "flag": "", "time": ""}],
                        ensure_ascii=False)
                items.append((crit.id, v, comment))

            total, _ = calculate_scores(items, cl)
            if total_file is not None and abs(total - float(total_file)) > 0.15:
                print(f"  ОШИБКА {name}: сделка {deal_id}, {operator}: "
                      f"пересчёт {total} ≠ файл {total_file}")
                errors += 1

            key = (deal_id, cl.id, call_dt.date() if call_dt else None, operator, round(total, 1))
            if key in existing:
                dup += 1
                continue
            existing.add(key)

            ev_user = users.get(evaluator)
            if not ev_user:
                print(f"  ОШИБКА {name}: оценщик «{evaluator}» не найден среди пользователей")
                errors += 1
            ev = Evaluation(
                checklist_id=cl.id,
                deal_id=deal_id,
                deal_url=f"https://entera.bitrix24.ru/crm/deal/details/{deal_id}/" if deal_id else None,
                operator_name=operator,
                eval_date=call_dt,
                week_num=call_dt.isocalendar()[1] if call_dt else None,
                week_year=call_dt.year if call_dt else None,
                month=MONTH_NAMES[call_dt.month - 1] if call_dt else None,
                department=dept if dept and dept != "—" else None,
                stage=_STAGES.get(stage_label, stage_label),
                total_score=total,
                evaluator_id=ev_user.id if ev_user else None,
                general_comment=(str(general).strip() or None) if general else None,
                status="published",
                created_at=eval_dt or datetime.utcnow(),
                is_calibration=False,
            )
            ev.items = [EvaluationItem(criterion_id=c, value=v, comment=cm) for c, v, cm in items]
            to_add.append(ev)
            added += 1
        print(f"{name}: чек-лист {cl.id} «{cl.name}» — к импорту {added}, уже в базе {dup}")

    print(f"\nИтого к импорту: {len(to_add)} оценок, "
          f"{sum(len(e.items) for e in to_add)} ответов, ошибок: {errors}")
    if errors:
        raise SystemExit("Есть ошибки — импорт не выполнен")
    if args.dry_run:
        print("--dry-run: в базу ничего не записано")
        return
    db.add_all(to_add)
    db.commit()
    print("Записано в базу.")


if __name__ == "__main__":
    main()
